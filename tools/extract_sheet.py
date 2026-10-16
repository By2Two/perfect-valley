import hashlib
import json
import os
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_JSON = os.path.join(ROOT, 'src', 'data', 'sheet.json')
OUTPUT_ICONS = os.path.join(ROOT, 'public', 'icons')
DEFAULT_WORKBOOK = 'Stardew_Valley_Completionist_Checklist__Vanilla_.xlsx'
ICON_HASH_LENGTH = 12
NAMESPACES = {
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'm': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
}
COLLECTION_SHEETS = {
    'shipping': ('Shipping', True),
    'cooking': ('Cooking', True),
    'crafting': ('Crafting', True),
    'fishing': ('Fishing', True),
    'museum': ('Museum', True),
    'gear': ('Gear', True),
    'tailoring': ('Tailoring', True),
    'powers': ('Special Powers', True),
    'statues': ('Misc.', True),
    'bestiary': ('Monsters', False),
}
GENERIC_LABELS = {'Product', 'Item', 'Monster', 'Recipe - Page 1', 'Recipe - Page 2', 'Recipe - Page 3', 'Cooking Recipes'}
SKILL_NAMES = ('Farming', 'Mining', 'Foraging', 'Fishing', 'Combat')
BACHELOR_STATUSES = ('Bachelor', 'Bachelorette')


def slugify(text):
    return re.sub(r'[^a-z0-9]+', '-', text.lower()).strip('-')


def read_relationships(archive, path):
    try:
        return {rel.get('Id'): rel.get('Target') for rel in ET.fromstring(archive.read(path))}
    except KeyError:
        return {}


def resolve_path(base, target):
    return os.path.normpath(os.path.join(os.path.dirname(base), target)).replace('\\', '/')


def relationships_path(path):
    return f'{os.path.dirname(path)}/_rels/{os.path.basename(path)}.rels'


def read_drawing_images(archive, drawing):
    drawing_relationships = read_relationships(archive, relationships_path(drawing))
    images = []
    for anchor in ET.fromstring(archive.read(drawing)):
        origin = anchor.find('xdr:from', NAMESPACES)
        blip = anchor.find('.//a:blip', NAMESPACES)
        if origin is None or blip is None:
            continue
        image = resolve_path(drawing, drawing_relationships[blip.get(f"{{{NAMESPACES['r']}}}embed")])
        row = int(origin.find('xdr:row', NAMESPACES).text) + 1
        column = int(origin.find('xdr:col', NAMESPACES).text) + 1
        images.append((row, column, archive.read(image)))
    return images


def read_image_anchors(path):
    archive = zipfile.ZipFile(path)
    workbook_relationships = read_relationships(archive, 'xl/_rels/workbook.xml.rels')
    anchors = {}
    for sheet in ET.fromstring(archive.read('xl/workbook.xml')).iter(f"{{{NAMESPACES['m']}}}sheet"):
        sheet_path = resolve_path('xl/workbook.xml', workbook_relationships[sheet.get(f"{{{NAMESPACES['r']}}}id")])
        images = []
        for target in read_relationships(archive, relationships_path(sheet_path)).values():
            if 'drawings/drawing' in target:
                images.extend(read_drawing_images(archive, resolve_path(sheet_path, target)))
        anchors[sheet.get('name')] = images
    return anchors


def save_icon(data, saved_hashes):
    icon_hash = hashlib.sha1(data).hexdigest()[:ICON_HASH_LENGTH]
    if icon_hash not in saved_hashes:
        with open(os.path.join(OUTPUT_ICONS, f'{icon_hash}.png'), 'wb') as icon_file:
            icon_file.write(data)
        saved_hashes.add(icon_hash)
    return icon_hash


def cell_text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip() or None


def find_header_row(sheet, row, column, image_cells):
    for candidate in range(row - 1, 0, -1):
        value = sheet.cell(candidate, column + 1).value
        check = sheet.cell(candidate, column - 1).value if column > 1 else None
        if isinstance(value, str) and not isinstance(check, bool) and (candidate, column) not in image_cells:
            return candidate
    return None


def has_group_row_above(sheet, header_row, column, image_cells):
    if not header_row or header_row <= 1 or (header_row - 1, column) in image_cells:
        return False
    check_above = sheet.cell(header_row - 1, column - 1).value if column > 1 else None
    value_above = sheet.cell(header_row - 1, column + 1).value
    return not isinstance(check_above, bool) and not isinstance(value_above, (int, float))


def apply_group_label(sheet, header_row, column, section):
    for current in range(column + 1, 0, -1):
        group = cell_text(sheet.cell(header_row - 1, current).value)
        if group and not group.startswith('=') and not group.replace('.', '').isdigit():
            if section in GENERIC_LABELS or section is None:
                return group
            return section if 'Progress' in group else f'{group} · {section}'
    return section


def find_section(sheet, row, column, image_cells):
    header_row = find_header_row(sheet, row, column, image_cells)
    section = cell_text(sheet.cell(header_row, column + 1).value) if header_row else None
    if has_group_row_above(sheet, header_row, column, image_cells):
        section = apply_group_label(sheet, header_row, column, section)
    is_generic = section in GENERIC_LABELS or (section and section.startswith('Recipe - Page')) or sheet.title == 'Crafting'
    return (None if is_generic else section), header_row


def read_details(sheet, row, column, header_row):
    details = []
    if not header_row:
        return details
    current = column + 2
    while current <= sheet.max_column:
        label = cell_text(sheet.cell(header_row, current).value)
        if not label or isinstance(sheet.cell(row, current).value, bool):
            break
        value = cell_text(sheet.cell(row, current).value)
        if value:
            details.append([label, value])
        current += 1
    return details


def extract_collection(sheet, commented_sheet, anchors, requires_check, saved_hashes):
    image_cells = {(row, column) for row, column, _ in anchors}
    items = []
    id_counts = {}
    for row, column, data in sorted(anchors):
        name = cell_text(sheet.cell(row, column + 1).value)
        check = sheet.cell(row, column - 1).value if column > 1 else None
        if not name or (requires_check and not isinstance(check, bool)):
            continue
        section, header_row = find_section(sheet, row, column, image_cells)
        note = commented_sheet.cell(row, column + 1).comment
        base_id = slugify(name)
        id_counts[base_id] = id_counts.get(base_id, 0) + 1
        item = {
            'id': base_id if id_counts[base_id] == 1 else f'{base_id}-{id_counts[base_id]}',
            'name': name,
            'icon': save_icon(data, saved_hashes),
        }
        details = read_details(sheet, row, column, header_row)
        if section:
            item['section'] = section
        if details:
            item['details'] = details
        if note and note.text.strip():
            item['note'] = note.text.strip()
        if check is True:
            item['fromSpreadsheet'] = 1
        items.append(item)
    return items


def extract_villagers(sheet, commented_sheet, anchors, saved_hashes):
    villagers = []
    for row, column, data in sorted(anchors):
        raw = cell_text(sheet.cell(row, column + 1).value)
        if not raw:
            continue
        name, _, birthday = [part.strip() for part in raw.partition('-')]
        villager = {'name': name, 'birthday': birthday, 'icon': save_icon(data, saved_hashes)}
        villager['isMarriageCandidate'] = cell_text(sheet.cell(row + 1, column).value) in BACHELOR_STATUSES
        for row_offset, key in ((2, 'loves'), (3, 'likes')):
            comment = commented_sheet.cell(row + row_offset, column).comment
            if comment:
                villager[key] = [line.strip() for line in comment.text.splitlines() if line.strip()]
        if sheet.cell(row, column - 1).value is True:
            villager['fromSpreadsheet'] = 1
        villagers.append(villager)
    return villagers


def extract_walnuts(sheet):
    walnuts = []
    for first_column in (1, 5):
        section = None
        for row in range(5, sheet.max_row + 1):
            checked = sheet.cell(row, first_column).value
            kind = cell_text(sheet.cell(row, first_column + 2).value)
            if sheet.cell(row, first_column + 1).value == '#':
                section = (kind or '').rstrip(' -').strip()
                continue
            if not isinstance(checked, bool) or not kind:
                continue
            match = re.search(r'\((\d+)\)', kind)
            source = cell_text(sheet.cell(row, first_column + 3).value)
            walnut = {
                'id': f'{slugify(section or "x")}-{len(walnuts) + 1}',
                'kind': re.sub(r'\(\d+\)', '', kind).strip(),
                'points': int(match.group(1)) if match else 1,
                'section': section,
            }
            if source:
                walnut['source'] = source
            if checked:
                walnut['fromSpreadsheet'] = 1
            walnuts.append(walnut)
    return walnuts


def extract_monster_goals(sheet):
    goals = []
    for row in range(5, 40):
        name = cell_text(sheet.cell(row, 10).value)
        goal = sheet.cell(row, 11).value
        if not name or not isinstance(goal, (int, float)):
            continue
        entry = {
            'id': slugify(name),
            'name': name,
            'goal': int(goal),
            'reward': cell_text(sheet.cell(row, 12).value),
            'description': cell_text(sheet.cell(row, 13).value),
        }
        if sheet.cell(row, 9).value is True:
            entry['fromSpreadsheet'] = 1
        goals.append(entry)
    return goals


def extract_perfection(sheet):
    stardrops = []
    obelisks = []
    for row in range(2, 9):
        name = cell_text(sheet.cell(row, 7).value)
        if name:
            stardrop = {'id': slugify(name), 'name': name, 'source': cell_text(sheet.cell(row, 8).value)}
            if sheet.cell(row, 6).value is True:
                stardrop['fromSpreadsheet'] = 1
            stardrops.append(stardrop)
    for row in range(11, 15):
        name = cell_text(sheet.cell(row, 7).value)
        if name:
            obelisk = {'id': slugify(name), 'name': name}
            if sheet.cell(row, 6).value is True:
                obelisk['fromSpreadsheet'] = 1
            obelisks.append(obelisk)
    has_clock = sheet.cell(4, 2).value == 'Yes'
    return stardrops, obelisks, has_clock


def extract_skills(sheet):
    skills = {}
    for row in range(1, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(row, column).value
            if value in SKILL_NAMES and isinstance(sheet.cell(row, column + 1).value, bool):
                skills[value.lower()] = sum(1 for current in range(column + 1, column + 14) if sheet.cell(row, current).value is True)
    return skills


def main(path):
    os.makedirs(OUTPUT_ICONS, exist_ok=True)
    os.makedirs(os.path.dirname(OUTPUT_JSON), exist_ok=True)
    anchors = read_image_anchors(path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    commented_workbook = openpyxl.load_workbook(path)
    saved_hashes = set()
    collections = {}
    for collection_id, (sheet_name, requires_check) in COLLECTION_SHEETS.items():
        collections[collection_id] = extract_collection(
            workbook[sheet_name], commented_workbook[sheet_name], anchors[sheet_name], requires_check, saved_hashes
        )
    stardrops, obelisks, has_clock = extract_perfection(workbook['Perfection'])
    data = {
        'collections': collections,
        'villagers': extract_villagers(workbook['Villagers'], commented_workbook['Villagers'], anchors['Villagers'], saved_hashes),
        'walnuts': extract_walnuts(workbook['Golden Walnuts']),
        'monsterGoals': extract_monster_goals(workbook['Monsters']),
        'stardrops': stardrops,
        'obelisks': obelisks,
        'fromSpreadsheet': {'clock': has_clock, 'skills': extract_skills(workbook['Farmer Skills'])},
    }
    with open(OUTPUT_JSON, 'w', encoding='utf-8') as output:
        json.dump(data, output, ensure_ascii=False, separators=(',', ':'))
    for collection_id, items in collections.items():
        print(f'{collection_id:10} {len(items):4} items')
    print(
        'villagers', len(data['villagers']),
        '| walnuts', sum(walnut['points'] for walnut in data['walnuts']),
        '| monster goals', len(data['monsterGoals']),
        '| icons', len(saved_hashes),
    )


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_WORKBOOK)
